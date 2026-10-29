"""Generate test fixtures (python-pptx + Pillow; LibreOffice + PyMuPDF optional for the visual reference)."""
import pathlib, shutil, subprocess
HERE = pathlib.Path(__file__).resolve().parent; FIX = HERE / 'fixtures'; FIX.mkdir(exist_ok=True)
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from PIL import Image, ImageDraw
im = Image.new('RGB', (640, 360), (30, 60, 120)); d = ImageDraw.Draw(im)
for i in range(0, 640, 40): d.line([(i, 0), (640 - i, 360)], fill=(255, 200, 0), width=3)
d.ellipse([220, 80, 420, 280], outline=(255, 255, 255), width=8); im.save(FIX / 'pic.png')
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[0]); s.shapes.title.text = 'U2B24-E DFA 内部培训 — ChipBench 测试'; s.placeholders[1].text = 'Wenhao Zhao · Application Engineering\nVersion 3.1'
s = prs.slides.add_slide(prs.slide_layouts[1]); s.shapes.title.text = '要点 / Key points'; tf = s.placeholders[1].text_frame; tf.text = '第一级：DFA 由 Base Block 组成'
for txt, lv in [('第二级：MAC / AddSub / SWC / Combine', 1), ('Level 2: CAN FD controller with message RAM', 1), ('第三级细节 detail text', 2), ('Back to level 1: clock tree & PLL', 0)]:
    p = tf.add_paragraph(); p.text = txt; p.level = lv
r = tf.paragraphs[0].runs[0]; r.font.bold = True; r.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
s = prs.slides.add_slide(prs.slide_layouts[6]); tb = s.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(0.9)); tb.text_frame.text = '自由形状 / Shapes & connectors'; tb.text_frame.paragraphs[0].runs[0].font.size = Pt(32); tb.text_frame.paragraphs[0].runs[0].font.bold = True
def box(x, y, w, h, txt, shape=MSO_SHAPE.ROUNDED_RECTANGLE, rgb=(0x3B, 0x66, 0xF0)):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(*rgb); sh.line.color.rgb = RGBColor(0x1E, 0x3A, 0x8A); sh.line.width = Pt(1.5)
    sh.text_frame.text = txt; p = sh.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; p.runs[0].font.size = Pt(16); p.runs[0].font.color.rgb = RGBColor(255, 255, 255); return sh
box(0.8, 2, 3, 1.3, 'MCU Core'); box(5.2, 2, 3, 1.3, 'DFA IP', rgb=(0x1F, 0x8A, 0x4C)); box(9.6, 2, 3, 1.3, 'Memory', MSO_SHAPE.RECTANGLE, (0xB2, 0x6B, 0x00))
for x1, x2 in [(3.8, 5.2), (8.2, 9.6)]:
    cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(2.65), Inches(x2), Inches(2.65)); cn.line.width = Pt(2.5); cn.line.color.rgb = RGBColor(0x44, 0x44, 0x44)
box(1, 4.2, 2.4, 1.4, 'Decision', MSO_SHAPE.DIAMOND, (0x8B, 0x5C, 0xF6)); box(4.2, 4.2, 2.4, 1.4, 'Process', MSO_SHAPE.RIGHT_ARROW, (0xD9, 0x48, 0x0F)); box(7.4, 4.2, 2.2, 1.4, 'Cloud', MSO_SHAPE.OVAL, (0x0E, 0x9A, 0xA7)); box(10.2, 4.2, 2.6, 1.4, 'Rotated', MSO_SHAPE.RECTANGLE, (0xC2, 0x25, 0x5C)).rotation = 12
s = prs.slides.add_slide(prs.slide_layouts[5]); s.shapes.title.text = '图片与表格 / Picture & table'; s.shapes.add_picture(str(FIX / 'pic.png'), Inches(0.7), Inches(1.7), Inches(5.6))
tbl = s.shapes.add_table(5, 3, Inches(6.9), Inches(1.7), Inches(5.8), Inches(3)).table
for i, row in enumerate([('Item', 'Value', 'Note'), ('fclk', '240 MHz', 'PLL0 / 2'), ('CAN', '500 kbps', 'sample 80%'), ('UART', '115200', 'err 0.2%'), ('ADC', '12 bit', '—')]):
    for j, v in enumerate(row): tbl.cell(i, j).text = v
s = prs.slides.add_slide(prs.slide_layouts[6]); g = s.shapes.add_group_shape()
for i in range(3):
    sh = g.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1 + i * 1.4), Inches(1.2), Inches(1.2), Inches(1.2)); sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(60 * i + 40, 100, 200 - 40 * i); sh.text_frame.text = str(i + 1)
cd = CategoryChartData(); cd.categories = ['A', 'B', 'C']; cd.add_series('S1', (3, 5, 2)); s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(3), Inches(6), Inches(3.5), cd)
t = s.shapes.add_textbox(Inches(7.5), Inches(1.2), Inches(5), Inches(3)); t.text_frame.word_wrap = True; t.text_frame.text = 'Bold, italic, underline, 上标'; p = t.text_frame.paragraphs[0]; p.runs[0].font.size = Pt(20)
r2 = p.add_run(); r2.text = ' and colored'; r2.font.color.rgb = RGBColor(0x1F, 0x8A, 0x4C); r2.font.bold = True; r2.font.size = Pt(20)
s.notes_slide.notes_text_frame.text = '这是演讲者备注 / speaker notes'; prs.save(str(FIX / 't1.pptx'))
big = Image.new('RGB', (1600, 900)); px = big.load()
for y in range(900):
    for x in range(0, 1600, 4):
        v = (x * 255 // 1600, y * 255 // 900, (x + y) % 255)
        for k in range(4): px[x + k, y] = v
big.save(FIX / 'big.jpg', quality=90)
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
for i in range(120):
    s = prs.slides.add_slide(prs.slide_layouts[5]); s.shapes.title.text = 'Slide %d — 性能测试 / performance' % (i + 1); s.shapes.add_picture(str(FIX / 'big.jpg'), Inches(0.6), Inches(1.6), Inches(6.2))
    t = s.shapes.add_table(6, 4, Inches(7.2), Inches(1.6), Inches(5.6), Inches(3)).table
    for rr in range(6):
        for c in range(4): t.cell(rr, c).text = '%d-%d' % (rr, c)
    tb = s.shapes.add_textbox(Inches(0.6), Inches(5.6), Inches(12), Inches(1.2)); tb.text_frame.text = '要点 %d: ' % i + 'DFA Base Block, MAC/AddSub/SWC. ' * 3
prs.save(str(FIX / 'big120.pptx'))
W = FIX / 'web'; (W / 'assets').mkdir(parents=True, exist_ok=True); Image.new('RGB', (40, 40), (220, 40, 40)).save(W / 'assets' / 'red.png')
(W / 'assets' / 'style.css').write_text("body{font-family:sans-serif} #box{width:80px;height:40px;background:url('red.png') no-repeat;border:2px solid #333} .big{color:rgb(0,128,0)}", encoding='utf-8')
(W / 'assets' / 'app.js').write_text("document.getElementById('out').textContent='script-ran:'+(typeof localStorage.getItem('k'))+':'+localStorage.getItem('k'); localStorage.setItem('k','v1'); document.getElementById('cnt').textContent=(Number(localStorage.getItem('n')||0)+1); localStorage.setItem('n',Number(localStorage.getItem('n')||0)+1);", encoding='utf-8')
(W / 'page.html').write_text("""<!doctype html><html><head><meta charset=utf-8><title>Local app</title>
<link rel=stylesheet href="assets/style.css"><meta http-equiv="Content-Security-Policy" content="script-src 'none'"></head>
<body><h1 class=big>Hello</h1><div id=box></div><img id=im src="assets/red%20.png" onerror="this.dataset.err=1"><img id=im2 src="assets/red.png">
<p id=out>no</p><p id=cnt></p><a id=lnk href="https://example.com/x">ext</a>
<script src="assets/app.js"></script><script>document.title='ok';</script></body></html>""", encoding='utf-8')
soffice = shutil.which('soffice') or shutil.which('libreoffice')
if soffice:
    try:
        subprocess.run([soffice, '--headless', '--convert-to', 'pdf', '--outdir', str(FIX), str(FIX / 't1.pptx')], capture_output=True, timeout=180)
        import fitz
        for i, pg in enumerate(fitz.open(str(FIX / 't1.pdf'))): pg.get_pixmap(matrix=fitz.Matrix(1.0, 1.0)).save(str(FIX / ('ref_%d.png' % (i + 1))))
        print('LibreOffice reference renders written')
    except Exception as e: print('reference renders skipped:', e)
print('fixtures ready in', FIX)
